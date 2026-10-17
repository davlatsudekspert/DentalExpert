"""Educational decks: validate drafts, optionally (re)generate them with the Claude API, build PPTX, render slide images.

  python3 tools/decks.py validate
  python3 tools/decks.py build --review        # same for decks/generated into decks/review (nothing public)
  python3 tools/decks.py build                 # PPTX -> decks/out/*.pptx (never served), slides -> public/decks/<id>/NN.webp + manifest.json
  ANTHROPIC_API_KEY=... python3 tools/decks.py generate [--ids tooth-layers,...]   # writes decks/generated/<id>.json for review

Slide images are the only thing the app serves. PPTX files stay in decks/out (git-ignored) and CI artifacts.
Generated drafts are NOT used by `build` until a person copies them over decks/drafts after reviewing them.
"""
from pathlib import Path
import json, os, re, shutil, subprocess, sys, tempfile
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DRAFTS, GENERATED, OUT = ROOT / 'decks/drafts', ROOT / 'decks/generated', ROOT / 'decks/out'
PUBLIC = ROOT / 'public/decks'
CATALOG = json.loads((ROOT / 'decks/catalog.json').read_text())
LAYOUTS = {'title', 'image-left', 'image-right', 'bullets', 'compare', 'closing'}
NAVY, TEAL, ICE, WHITE, GREY, MINT = (RGBColor.from_string(x) for x in ('10254E', '17A6B7', 'F0FAFF', 'FFFFFF', '5C7590', 'E8F5EE'))
FONT = 'Liberation Sans'
FOOTER = 'Dental Expert · o‘quv namuna · mutaxassis tasdig‘idan o‘tmagan'
W, H = Inches(13.333), Inches(7.5)

# ---------------------------------------------------------------- validation
BAD = [re.compile(p, re.I) for p in (r'https?://', r'www\.', r'\(\s*\d{4}\s*\)', r'\b(p|pp)\.\s*\d', r'\b\d+\s*[-–]\s*bet\b', r'\bbet\s*\d+', r'\bsahifa\s*\d+', r'\bISBN\b', r'\bdoi\b')]


def validate(deck: dict) -> list[str]:
    errs = []
    def text_ok(where, s):
        for p in BAD:
            if p.search(s):
                errs.append(f'{where}: citation/link-like text "{s[:50]}"')
    for key in ('id', 'title', 'subtitle', 'audience'):
        if not isinstance(deck.get(key), str) or not deck[key].strip():
            errs.append(f'missing {key}')
    slides = deck.get('slides') or []
    if not 6 <= len(slides) <= 12:
        errs.append(f'slide count {len(slides)} not in 6..12')
    for i, s in enumerate(slides, 1):
        w = f'slide {i}'
        if s.get('layout') not in LAYOUTS:
            errs.append(f'{w}: bad layout {s.get("layout")}')
        if not s.get('title') or len(s['title']) > 70:
            errs.append(f'{w}: title missing or >70 chars')
        text_ok(w, s.get('title', ''))
        for b in s.get('bullets', []):
            if len(b) > 170:
                errs.append(f'{w}: bullet >170 chars')
            text_ok(w, b)
        if len(s.get('bullets', [])) > 5:
            errs.append(f'{w}: more than 5 bullets')
        for k in ('image', 'image2'):
            if s.get(k) and s[k] not in CATALOG:
                errs.append(f'{w}: image "{s[k]}" not in catalog')
        if s.get('layout') in ('image-left', 'image-right') and not s.get('image'):
            errs.append(f'{w}: image layout without image')
        if s.get('layout') == 'compare' and not (s.get('image') and s.get('image2')):
            errs.append(f'{w}: compare needs image and image2')
    quiz = deck.get('quiz') or []
    if not 3 <= len(quiz) <= 8:
        errs.append(f'quiz size {len(quiz)} not in 3..8')
    for i, q in enumerate(quiz, 1):
        if not (isinstance(q.get('options'), list) and len(q['options']) == 3 and isinstance(q.get('correct'), int) and 0 <= q['correct'] < 3 and q.get('q') and q.get('explanation')):
            errs.append(f'quiz {i}: needs q, 3 options, correct 0..2, explanation')
        text_ok(f'quiz {i}', ' '.join([q.get('q', ''), q.get('explanation', '')] + list(q.get('options', []))))
    return errs

# ---------------------------------------------------------------- images
_cache: dict[str, Path] = {}


def image_png(key: str) -> tuple[Path, tuple[int, int]]:
    if key not in _cache:
        im = Image.open(ROOT / CATALOG[key]['file']).convert('RGBA')
        bb = im.split()[3].getbbox()
        if bb:
            im = im.crop(bb)
        p = Path(tempfile.gettempdir()) / f'deck_{key}.png'
        im.save(p)
        _cache[key] = p
        _cache[key + ':size'] = im.size  # type: ignore[assignment]
    return _cache[key], _cache[key + ':size']  # type: ignore[return-value]

# ---------------------------------------------------------------- pptx helpers


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, line=None):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line: s.line.color.rgb = line; s.line.width = Pt(1)
    else: s.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, s, size, color=NAVY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = spacing
    r = p.add_run(); r.text = s; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    return tb


def picture_card(slide, key, x, y, w, h, caption=None):
    rect(slide, x, y, w, h, WHITE, radius=0.05, line=RGBColor.from_string('D3E6EE'))
    path, (iw, ih) = image_png(key)
    pad = Inches(0.25); cap = Inches(0.45) if caption else 0
    bw, bh = w - 2 * pad, h - 2 * pad - cap
    scale = min(bw / iw, bh / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + pad + (bh - ph) // 2, pw, ph)
    if caption:
        text(slide, x + pad, y + h - pad - Inches(0.3), w - 2 * pad, Inches(0.3), caption, 12, GREY, align=PP_ALIGN.CENTER)


def base(prs, n, total, title=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = ICE
    rect(slide, 0, 0, Inches(0.22), H, TEAL, shape=MSO_SHAPE.RECTANGLE)
    text(slide, Inches(0.7), Inches(7.0), Inches(9), Inches(0.3), FOOTER, 10, GREY)
    text(slide, Inches(11.6), Inches(7.0), Inches(1.2), Inches(0.3), f'{n} / {total}', 10, GREY, align=PP_ALIGN.RIGHT)
    if title:
        text(slide, Inches(0.7), Inches(0.55), Inches(11.9), Inches(1.1), title, 36, NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        rect(slide, Inches(0.7), Inches(1.7), Inches(1.1), Inches(0.07), TEAL, shape=MSO_SHAPE.RECTANGLE)
    return slide


def bullet_list(slide, items, x, y, w, size=22, gap=0.25):
    cur = y
    for it in items:
        lines = max(1, -(-len(it) * size // int(w / Inches(1) * 9.2 * 7.2 + 1)))
        h = Inches(0.42 * max(lines, 1) * size / 22 + 0.12)
        rect(slide, x, cur + Inches(0.1), Inches(0.14), Inches(0.14), TEAL, shape=MSO_SHAPE.OVAL)
        text(slide, x + Inches(0.38), cur, w - Inches(0.38), h, it, size, NAVY)
        cur += h + Inches(gap)


def build_deck(deck: dict, out: Path):
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    slides = deck['slides']; total = len(slides)
    for n, s in enumerate(slides, 1):
        lay, bl = s['layout'], s.get('bullets', [])
        if lay == 'title':
            sl = base(prs, n, total)
            text(sl, Inches(0.8), Inches(1.4), Inches(6.2), Inches(0.4), 'DENTAL EXPERT · DENTAL ACADEMY', 13, TEAL, bold=True)
            lines = max(1, -(-len(s['title']) // 15))
            title_h = Inches(0.82 * lines)
            text(sl, Inches(0.8), Inches(2.0), Inches(6.4), title_h, s['title'], 48, NAVY, bold=True, anchor=MSO_ANCHOR.TOP, spacing=0.95)
            sub_y = Inches(2.0) + title_h + Inches(0.35)
            text(sl, Inches(0.8), sub_y, Inches(6.2), Inches(0.9), ' · '.join(bl) or deck['subtitle'], 20, GREY)
            text(sl, Inches(0.8), sub_y + Inches(1.1), Inches(6.2), Inches(0.4), deck['audience'], 14, TEAL, bold=True)
            picture_card(sl, s['image'], Inches(7.6), Inches(0.9), Inches(5.1), Inches(5.7))
        elif lay in ('image-left', 'image-right'):
            sl = base(prs, n, total, s['title'])
            img_x, txt_x = (Inches(0.7), Inches(6.9)) if lay == 'image-left' else (Inches(7.6), Inches(0.7))
            picture_card(sl, s['image'], img_x, Inches(2.05), Inches(5.0), Inches(4.7), s.get('caption'))
            bullet_list(sl, bl, txt_x, Inches(2.25), Inches(5.7), size=22)
        elif lay == 'bullets':
            sl = base(prs, n, total, s['title'])
            k = max(len(bl), 1); gap = Inches(0.22)
            avail = Inches(4.75); rh = int((avail - gap * (k - 1)) / k); rh = min(rh, Inches(1.25))
            y = Inches(2.1)
            for i, b in enumerate(bl, 1):
                rect(sl, Inches(0.7), y, Inches(11.9), rh, WHITE, radius=0.12, line=RGBColor.from_string('D3E6EE'))
                circ = rect(sl, Inches(0.95), y + (rh - Inches(0.62)) // 2, Inches(0.62), Inches(0.62), TEAL, shape=MSO_SHAPE.OVAL)
                text(sl, Inches(0.95), y + (rh - Inches(0.62)) // 2, Inches(0.62), Inches(0.62), str(i), 20, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, Inches(1.9), y, Inches(10.4), rh, b, 22, NAVY, anchor=MSO_ANCHOR.MIDDLE)
                y += rh + gap
        elif lay == 'compare':
            sl = base(prs, n, total, s['title'])
            picture_card(sl, s['image'], Inches(0.7), Inches(2.05), Inches(5.8), Inches(4.75), s.get('caption'))
            picture_card(sl, s['image2'], Inches(6.8), Inches(2.05), Inches(5.8), Inches(4.75), s.get('caption2'))
        else:  # closing
            sl = base(prs, n, total, s['title'])
            k = max(len(bl), 1); gap = Inches(0.2); rh = min(Inches(1.0), int((Inches(4.7) - gap * (k - 1)) / k)); y = Inches(2.1)
            for b in bl:
                rect(sl, Inches(0.7), y, Inches(7.3), rh, MINT, radius=0.14)
                text(sl, Inches(1.0), y, Inches(0.5), rh, '✓', 24, TEAL, bold=True, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, Inches(1.65), y, Inches(6.1), rh, b, 21, NAVY, anchor=MSO_ANCHOR.MIDDLE)
                y += rh + gap
            if s.get('image'):
                picture_card(sl, s['image'], Inches(8.5), Inches(2.1), Inches(4.1), Inches(4.5))
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)

# ---------------------------------------------------------------- render


def render(pptx: Path, deck_id: str, public: Path = PUBLIC) -> int:
    with tempfile.TemporaryDirectory() as td:
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', td, str(pptx)], check=True, capture_output=True, timeout=240)
        pdf = Path(td) / (pptx.stem + '.pdf')
        subprocess.run(['pdftoppm', '-r', '96', '-png', str(pdf), str(Path(td) / 's')], check=True, capture_output=True, timeout=240)
        pngs = sorted(Path(td).glob('s-*.png'))
        dest = public / deck_id
        if dest.exists(): shutil.rmtree(dest)
        dest.mkdir(parents=True)
        for i, p in enumerate(pngs, 1):
            Image.open(p).convert('RGB').save(dest / f'{i:02d}.webp', 'WEBP', quality=82, method=6)
    return len(pngs)


def cmd_build(src: Path = DRAFTS, out: Path = OUT, public: Path = PUBLIC):
    manifest = []
    for f in sorted(src.glob('*.json')):
        deck = json.loads(f.read_text())
        errs = validate(deck)
        if errs:
            sys.exit(f'{f.name}: ' + '; '.join(errs))
        pptx = out / f'{deck["id"]}.pptx'
        build_deck(deck, pptx)
        n = render(pptx, deck['id'], public)
        assert n == len(deck['slides']), (deck['id'], n, len(deck['slides']))
        manifest.append({'id': deck['id'], 'title': deck['title'], 'subtitle': deck['subtitle'], 'audience': deck['audience'], 'slides': n, 'quiz': deck['quiz']})
        print('built', deck['id'], n, 'slides')
    public.mkdir(parents=True, exist_ok=True)
    (public / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=1) + '\n')


# ---------------------------------------------------------------- Claude API generation
PROMPT = """You write short, accurate educational slide decks for dental students in Uzbekistan. Write ALL text in Uzbek (Latin script, use ‘ and ’ for oʻ gʻ marks like the examples).

Topic: {title}
Subtitle: {subtitle}
Audience: {audience}
Outline hint (slide titles of the current draft, you may improve them): {outline}

Hard rules:
- Only well-established, textbook-level facts. If unsure, leave it out. No numbers, statistics, dates, percentages, drug names or doses unless universally standard (e.g. 32 permanent teeth).
- NEVER invent or imply sources: no book titles, authors, page numbers, years, URLs, DOIs. Do not write "according to ...".
- No diagnosis or treatment instructions. Where relevant say that diagnosis and treatment belong to a dentist.
- Do not claim anything about the number of roots or canals shown in the pictures.
- 8 slides. Layouts: title (first), then a mix of image-left, image-right, bullets, compare, and closing (last). Titles max 70 chars, bullets max 160 chars, max 4 bullets per slide.
- Use ONLY these image ids, and only when the picture really fits: {images}
- title and closing slides need an image; image-left/right need image; compare needs image and image2 (two different ids) plus caption and caption2.
- Also write 5 multiple-choice quiz questions, exactly 3 options each, "correct" is the 0-based index, with a one-sentence explanation. Vary the position of the correct option.

Return ONLY JSON (no markdown) with this exact shape:
{{"id": "{id}", "title": str, "subtitle": str, "audience": str,
 "slides": [{{"layout": str, "title": str, "bullets": [str], "image": str, "caption": str, "image2": str, "caption2": str}}],
 "quiz": [{{"q": str, "options": [str, str, str], "correct": int, "explanation": str}}]}}
Omit keys that a slide does not use."""


def cmd_generate(ids):
    import anthropic
    client = anthropic.Anthropic()
    model = os.environ.get('DECK_MODEL', 'claude-sonnet-5-5')
    GENERATED.mkdir(parents=True, exist_ok=True)
    images = '; '.join(f'{k} = {v["about"]}' for k, v in CATALOG.items())
    for f in sorted(DRAFTS.glob('*.json')):
        draft = json.loads(f.read_text())
        if ids and draft['id'] not in ids:
            continue
        prompt = PROMPT.format(title=draft['title'], subtitle=draft['subtitle'], audience=draft['audience'], id=draft['id'],
                               outline=' | '.join(s['title'] for s in draft['slides']), images=images)
        last = []
        for attempt in range(3):
            msg = client.messages.create(model=model, max_tokens=6000, messages=[{'role': 'user', 'content': prompt + ('\n\nFix these problems from your previous answer: ' + '; '.join(last) if last else '')}])
            raw = ''.join(b.text for b in msg.content if b.type == 'text').strip()
            raw = re.sub(r'^```(?:json)?|```$', '', raw, flags=re.M).strip()
            try:
                deck = json.loads(raw)
                deck['id'] = draft['id']
                last = validate(deck)
            except Exception as e:  # malformed JSON
                last = [f'invalid JSON: {e}']
            if not last:
                (GENERATED / f'{draft["id"]}.json').write_text(json.dumps(deck, ensure_ascii=False, indent=1) + '\n')
                print('generated', draft['id'])
                break
        else:
            print('FAILED', draft['id'], last)
            sys.exit(1)


if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else 'build'
    if cmd == 'validate':
        bad = 0
        for f in sorted(DRAFTS.glob('*.json')):
            e = validate(json.loads(f.read_text()))
            print(f.name, 'OK' if not e else e); bad += bool(e)
        sys.exit(1 if bad else 0)
    elif cmd == 'build':
        if '--review' in sys.argv:  # build the Claude-generated drafts into decks/review without touching public/
            cmd_build(GENERATED, ROOT / 'decks/review/pptx', ROOT / 'decks/review/slides')
        else:
            cmd_build()
    elif cmd == 'generate':
        ids = None
        if '--ids' in sys.argv:
            ids = set(sys.argv[sys.argv.index('--ids') + 1].split(','))
        cmd_generate(ids)
    else:
        sys.exit(__doc__)
