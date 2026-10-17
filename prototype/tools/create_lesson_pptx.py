"""Build the original Uzbek visual anatomy demo deck from the user's tooth image."""
from pathlib import Path
from tempfile import TemporaryDirectory

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'assets/academy/tooth-layers-demo-uz.pptx'
SOURCE=ROOT/'assets/tooth3d/premium-reference-cutaway-ivory.webp'
NAVY=RGBColor(10,32,71)
TEAL=RGBColor(0,145,164)
PALE=RGBColor(232,248,252)
GREY=RGBColor(90,119,139)

def text(slide,body,x,y,w,h,size=24,bold=False,color=NAVY):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    frame=shape.text_frame
    frame.word_wrap=True
    paragraph=frame.paragraphs[0]
    paragraph.text=body
    paragraph.font.name='Aptos'
    paragraph.font.size=Pt(size)
    paragraph.font.bold=bold
    paragraph.font.color.rgb=color
    return shape

def base(prs,eyebrow,title):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb=RGBColor(247,252,254)
    text(slide,eyebrow,0.65,0.32,11.8,0.35,13,True,TEAL)
    text(slide,title,0.65,0.78,11.7,0.8,29,True)
    text(slide,'DENTAL EXPERT  ·  VIZUAL DEMO',0.65,7.03,7.8,0.3,10,False,GREY)
    return slide

def main():
    prs=Presentation()
    prs.slide_width=Inches(13.33)
    prs.slide_height=Inches(7.5)
    with TemporaryDirectory() as tmp:
        png=Path(tmp)/'cutaway.png'
        Image.open(SOURCE).convert('RGBA').save(png)

        slide=base(prs,'01  /  TISH MORFOLOGIYASI','Tish qatlamlari')
        text(slide,'Rasmli kuzatish va qisqa bilim tekshiruvi',0.75,2.18,6.6,0.7,22)
        text(slide,'Bemor ma’lumoti emas · klinik tashxis emas',0.75,3.12,6.5,0.4,15,False,GREY)
        slide.shapes.add_picture(str(png),Inches(7.0),Inches(1.35),height=Inches(5.45))

        slide=base(prs,'02  /  RASMGA QARANG','To‘rtta ko‘rinadigan qism')
        slide.shapes.add_picture(str(png),Inches(0.6),Inches(1.55),height=Inches(5.22))
        for i,(name,caption) in enumerate([('01  Emal','Tashqi toj qatlami'),('02  Dentin','Emal ostidagi qism'),('03  Pulpa','Markazdagi pushti qism'),('04  Ildiz','Pastga davom etgan qism')]):
            y=1.65+i*1.17
            text(slide,name,7.2,y,5.5,0.35,20,True,TEAL)
            text(slide,caption,7.2,y+0.42,5.5,0.38,15,False,GREY)

        slide=base(prs,'03  /  O‘ZINGIZ TEKSHIRING','Rasmda toping')
        for i,question in enumerate(['Tashqi oq qavatni ko‘rsating.','Markazdagi pushti joyni ko‘rsating.','Ildizlarni ko‘rsating.']):
            text(slide,f'{i+1:02d}',0.8,1.85+i*1.3,0.7,0.5,22,True,TEAL)
            text(slide,question,1.55,1.85+i*1.3,10.7,0.55,22)
        text(slide,'Javobni ilovadagi interaktiv testda tanlang.',0.8,6.13,11.5,0.42,16,False,GREY)

        slide=base(prs,'04  /  MANBA HOLATI','Nashrdan oldingi tekshiruv')
        text(slide,'Tasvir: loyiha egasi yuborgan tish kesimi asosidagi vizual.',0.8,1.78,11.8,0.75,21)
        text(slide,'Bu rasm klinik anatomiya sifatida tasdiqlanmagan. Darslik, nashr va aniq bet/bo‘lim manbalari hali tekshirilmagan; tasdiqlanmaguncha bu slaydlar faqat dizayn va o‘quv oqimi demosidir.',0.8,2.92,11.4,1.9,20,False,GREY)
        text(slide,'Ustoz: rasmni ko‘rsatib, talabadan sababli izoh va qayta ishlashni so‘rang.',0.8,5.5,11.5,0.7,17,True,TEAL)

    OUT.parent.mkdir(parents=True,exist_ok=True)
    prs.save(OUT)
    print(OUT.relative_to(ROOT),OUT.stat().st_size,'bytes')

if __name__=='__main__':main()
